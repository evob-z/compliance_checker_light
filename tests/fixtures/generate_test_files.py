"""
生成 tests/fixtures 下的最小二进制测试文件。

依赖（dev/可选）: reportlab、python-docx、Pillow。
从任意 cwd 运行均可：输出目录始终为本文件所在目录。
"""

from __future__ import annotations

import os
from pathlib import Path

from reportlab.pdfgen import canvas
from docx import Document
from PIL import Image, ImageDraw

FIXTURES_DIR = Path(__file__).resolve().parent
os.makedirs(FIXTURES_DIR, exist_ok=True)


def generate_pdf() -> Path:
    """生成用于测试完整性和基础解析的 PDF"""
    path = FIXTURES_DIR / "dummy_approval.pdf"
    c = canvas.Canvas(str(path))
    # 写入一些合规审查的关键词
    c.drawString(100, 750, "Document Name: Project Approval (dummy)")
    c.drawString(100, 730, "Project Name: AI System Construction")
    c.drawString(100, 710, "Issue Date: 2024-01-01")
    c.drawString(100, 690, "Status: Approved")
    c.save()
    print(f"✅ 生成 PDF: {path}")
    return path


def generate_docx() -> Path:
    """生成用于测试时效性检查的 DOCX"""
    path = FIXTURES_DIR / "dummy_permit.docx"
    doc = Document()
    doc.add_heading("Construction Permit (Test)", 0)
    doc.add_paragraph("This is a dummy construction permit for testing purposes.")
    doc.add_paragraph("Valid Until: 2025-12-31")  # 用于测试时效性
    doc.save(str(path))
    print(f"✅ 生成 DOCX: {path}")
    return path


def generate_image() -> Path:
    """生成用于测试视觉检查器（印章/签名）的假图片"""
    path = FIXTURES_DIR / "dummy_seal.jpg"
    # 创建一个白色背景图片
    img = Image.new("RGB", (300, 200), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    # 画一个红色的圆圈模拟印章
    d.ellipse([100, 50, 200, 150], outline=(255, 0, 0), width=5)
    d.text((120, 95), "FAKE SEAL", fill=(255, 0, 0))
    d.text((10, 10), "法人签名: 张三 (测试)", fill=(0, 0, 0))
    img.save(path)
    print(f"✅ 生成 JPG: {path}")
    return path


if __name__ == "__main__":
    generate_pdf()
    generate_docx()
    generate_image()
