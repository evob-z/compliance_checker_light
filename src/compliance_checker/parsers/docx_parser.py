"""
Word 文档解析器 - 支持 .docx 格式
"""
import logging
from pathlib import Path
from typing import Optional, List
from datetime import datetime

from ..core.document import Document, DocumentType, PageContent, DocumentMetadata

logger = logging.getLogger(__name__)


class DocxParser:
    """
    Word文档解析器 (.docx)
    
    功能：
    - 提取文档文本内容
    - 提取文档元数据
    - 提取页数（估算）
    """
    
    def __init__(self):
        """初始化Word解析器"""
        self._docx_module = None
    
    def _get_docx_module(self):
        """延迟加载python-docx模块"""
        if self._docx_module is None:
            try:
                import docx
                self._docx_module = docx
            except ImportError:
                raise ImportError(
                    "python-docx 未安装，请运行: pip install python-docx"
                )
        return self._docx_module
    
    def parse(self, file_path: str) -> Document:
        """
        解析Word文档
        
        Args:
            file_path: Word文件路径
            
        Returns:
            Document对象
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
        
        if not path.suffix.lower() in ['.docx', '.doc']:
            raise ValueError(f"不支持的文件格式: {path.suffix}")
        
        if path.suffix.lower() == '.doc':
            logger.warning(".doc格式需要转换为.docx，建议先转换文件格式")
            # 可以尝试使用其他库转换，这里暂时不支持
            raise ValueError("暂不支持 .doc 格式，请先转换为 .docx")
        
        logger.info(f"开始解析Word文档: {file_path}")
        
        docx = self._get_docx_module()
        doc = docx.Document(file_path)
        
        try:
            # 提取元数据
            metadata = self._extract_metadata(doc, file_path)
            
            # 提取页面内容
            pages_content = self._extract_content(doc)
            
            # 估算页数（Word没有固定页数，按内容估算）
            estimated_pages = max(1, len(pages_content))
            
            document = Document(
                file_path=str(path.absolute()),
                file_name=path.name,
                file_type=DocumentType.DOCX,
                pages=estimated_pages,
                pages_content=pages_content,
                metadata=metadata
            )
            
            logger.info(f"Word文档解析完成: {path.name}")
            return document
            
        except Exception as e:
            logger.error(f"Word文档解析失败: {e}")
            raise
    
    def _extract_metadata(self, doc, file_path: str) -> DocumentMetadata:
        """提取Word文档元数据"""
        metadata = DocumentMetadata()
        
        try:
            # 提取核心属性
            core_props = doc.core_properties
            
            metadata.title = core_props.title
            metadata.author = core_props.author
            metadata.subject = core_props.subject
            metadata.creator = core_props.creator
            
            if core_props.created:
                metadata.creation_date = core_props.created
            if core_props.modified:
                metadata.modification_date = core_props.modified
                
        except Exception as e:
            logger.debug(f"提取元数据失败: {e}")
        
        return metadata
    
    def _extract_content(self, doc) -> List[PageContent]:
        """
        提取文档内容
        
        Word文档没有明确的页面概念，这里按段落组织内容
        """
        paragraphs = []
        
        # 提取所有段落文本
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)
        
        # 提取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        row_text.append(cell_text)
                if row_text:
                    paragraphs.append(" | ".join(row_text))
        
        # 将所有内容作为一个"页面"
        # 如果内容很多，可以按段落数分割成多个虚拟页面
        full_text = "\n".join(paragraphs)
        
        # 按大约3000字符分割为虚拟页面（便于后续处理）
        chars_per_page = 3000
        pages_content = []
        
        if len(full_text) <= chars_per_page:
            pages_content.append(PageContent(
                page_num=0,
                text=full_text,
                has_text_layer=True
            ))
        else:
            # 分割为多个虚拟页面
            page_num = 0
            for i in range(0, len(full_text), chars_per_page):
                page_text = full_text[i:i + chars_per_page]
                pages_content.append(PageContent(
                    page_num=page_num,
                    text=page_text,
                    has_text_layer=True
                ))
                page_num += 1
        
        return pages_content


def parse_docx(file_path: str) -> Document:
    """
    便捷函数：解析Word文档
    
    Args:
        file_path: Word文件路径
        
    Returns:
        Document对象
    """
    parser = DocxParser()
    return parser.parse(file_path)
